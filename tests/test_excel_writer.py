import io

from openpyxl import load_workbook

from server.excel_writer import generate_excel

HEADERS = ["Tanggal", "Keterangan", "CBG", "Mutasi", "DB/CR", "Saldo"]


def test_generate_excel_writes_headers_and_rows():
    data = [["01/02/2024", "TRANSFER", "0012", 1000.0, "CR", 5000.0]]
    out = io.BytesIO()
    generate_excel(data, {"bank": "BCA", "headers": HEADERS}, out)
    out.seek(0)
    wb = load_workbook(out)
    ws = wb.active
    assert ws.title == "Mutasi Rekening"
    assert [c.value for c in ws[1]] == HEADERS
    assert ws["A2"].value == "01/02/2024"
    assert ws["B2"].value == "TRANSFER"
    assert ws["D2"].value == 1000.0
    assert ws["E2"].value == "CR"
    assert ws["F2"].value == 5000.0


def test_generate_excel_applies_number_format():
    out = io.BytesIO()
    generate_excel([["x", 1234.5]], {"headers": ["A", "B"]}, out)
    out.seek(0)
    ws = load_workbook(out).active
    assert ws["B2"].number_format == "#,##0.00"


def test_generate_excel_empty_data():
    out = io.BytesIO()
    generate_excel([], {"headers": HEADERS}, out)
    out.seek(0)
    ws = load_workbook(out).active
    assert ws.max_row == 1
    assert [c.value for c in ws[1]] == HEADERS


def test_generate_excel_multiple_rows():
    data = [
        ["01/01", "A", "1", 100.0, "DB", 900.0],
        ["02/01", "B", "2", 200.0, "CR", 1100.0],
    ]
    out = io.BytesIO()
    generate_excel(data, {"headers": HEADERS}, out)
    out.seek(0)
    ws = load_workbook(out).active
    assert ws.max_row == 3
    assert ws["A3"].value == "02/01"
    assert ws["F3"].value == 1100.0
